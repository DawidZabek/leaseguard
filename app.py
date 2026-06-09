import os
import sys
import json
import uuid
import fitz
import docx
from fpdf import FPDF
from flask import Flask, render_template, request, jsonify, make_response
from werkzeug.utils import secure_filename
from dotenv import load_dotenv

load_dotenv()

sys.path.insert(0, os.path.dirname(__file__))

from rag.setup import setup_rag
from agents.extractor import run_extractor
from agents.advisor import run_advisor
from agents.photo import run_photo_analysis
from agents.protocol import run_protocol

app = Flask(__name__)
app.secret_key = os.getenv("SECRET_KEY", "leaseguard-dev-secret")
app.config["UPLOAD_FOLDER"] = os.path.join(os.path.dirname(__file__), "uploads")
app.config["MAX_CONTENT_LENGTH"] = 32 * 1024 * 1024

ALLOWED_IMAGES = {"png", "jpg", "jpeg", "webp"}
ALLOWED_DOCS = {"pdf", "docx", "doc", "txt"}

os.makedirs(app.config["UPLOAD_FOLDER"], exist_ok=True)


def _ext(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def extract_text_from_file(path: str, ext: str) -> str:
    if ext == "pdf":
        doc = fitz.open(path)
        return "\n".join(page.get_text() for page in doc)
    if ext in ("docx", "doc"):
        doc = docx.Document(path)
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    if ext == "txt":
        with open(path, encoding="utf-8", errors="replace") as f:
            return f.read()
    return ""


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/contract")
def contract_page():
    return render_template("contract.html")


@app.route("/protocol")
def protocol_page():
    return render_template("protocol.html")


@app.route("/api/upload-contract", methods=["POST"])
def upload_contract():
    if "file" not in request.files:
        return jsonify({"error": "Brak pliku"}), 400

    file = request.files["file"]
    if not file.filename:
        return jsonify({"error": "Nie wybrano pliku"}), 400

    ext = _ext(file.filename)
    if ext not in ALLOWED_DOCS:
        return jsonify({"error": f"Nieobsługiwany format. Akceptowane: PDF, DOCX, TXT"}), 400

    filename = f"{uuid.uuid4()}_{secure_filename(file.filename)}"
    path = os.path.join(app.config["UPLOAD_FOLDER"], filename)
    try:
        file.save(path)
        text = extract_text_from_file(path, ext)
    finally:
        try:
            os.remove(path)
        except OSError:
            pass

    if not text or len(text.strip()) < 100:
        return jsonify({"error": "Nie udało się wyodrębnić tekstu z pliku lub plik jest zbyt krótki"}), 400

    return jsonify({"text": text.strip(), "chars": len(text.strip())})


@app.route("/api/analyze-contract", methods=["POST"])
def analyze_contract():
    data = request.get_json()
    if not data or not data.get("contract_text"):
        return jsonify({"error": "Brak tekstu umowy"}), 400

    contract_text = data["contract_text"].strip()
    if len(contract_text) < 100:
        return jsonify({"error": "Tekst umowy jest zbyt krótki"}), 400

    try:
        clauses, clause_risks = run_extractor(contract_text)
        if not clauses:
            return jsonify({"error": "Nie wykryto żadnych klauzul w tekście"}), 400

        report = run_advisor(clause_risks)

        return jsonify({
            "clauses": [
                {
                    "clause_type": cr.clause.clause_type,
                    "content": cr.clause.content,
                    "raw_excerpt": cr.clause.raw_excerpt,
                    "article_reference": cr.clause.article_reference,
                    "status": cr.status,
                    "justification": cr.justification,
                    "legal_basis": cr.legal_basis,
                    "recommendation": cr.recommendation,
                }
                for cr in report.clauses
            ],
            "questions_for_landlord": report.questions_for_landlord,
            "overall_recommendation": report.overall_recommendation,
            "risk_summary": report.risk_summary,
        })
    except json.JSONDecodeError as e:
        return jsonify({"error": f"Błąd parsowania odpowiedzi AI: {str(e)}"}), 500
    except Exception as e:
        return jsonify({"error": f"Błąd analizy: {str(e)}"}), 500


@app.route("/api/analyze-photos", methods=["POST"])
def analyze_photos():
    if "photos" not in request.files:
        return jsonify({"error": "Brak zdjęć"}), 400

    files = request.files.getlist("photos")
    address = request.form.get("address", "")

    if not files or all(f.filename == "" for f in files):
        return jsonify({"error": "Nie wybrano żadnych plików"}), 400

    saved_paths = []
    try:
        for file in files:
            if file and _ext(file.filename) in ALLOWED_IMAGES:
                filename = f"{uuid.uuid4()}_{secure_filename(file.filename)}"
                path = os.path.join(app.config["UPLOAD_FOLDER"], filename)
                file.save(path)
                saved_paths.append(path)

        if not saved_paths:
            return jsonify({"error": "Żaden z przesłanych plików nie jest obsługiwanym formatem obrazu"}), 400

        rooms = [run_photo_analysis(path) for path in saved_paths]
        protocol = run_protocol(rooms, address)

        return jsonify({
            "rooms": [
                {
                    "room_name": r.room_name,
                    "defects": r.defects,
                    "general_condition": r.general_condition,
                    "recommendations": r.recommendations,
                    "photo_description": r.photo_description,
                }
                for r in protocol.rooms
            ],
            "protocol_text": protocol.protocol_text,
            "total_defects": protocol.total_defects,
            "property_address": protocol.property_address,
        })
    except Exception as e:
        return jsonify({"error": f"Błąd analizy zdjęć: {str(e)}"}), 500
    finally:
        for path in saved_paths:
            try:
                os.remove(path)
            except OSError:
                pass


DEJAVU_FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
DEJAVU_BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"


def _build_protocol_pdf(text: str) -> bytes:
    def wrap_words(line: str, max_chars: int = 80) -> str:
        parts = []
        for word in line.split(" "):
            while len(word) > max_chars:
                parts.append(word[:max_chars])
                word = word[max_chars:]
            parts.append(word)
        return " ".join(parts)

    # Marginesy przed add_page — to jest kluczowe
    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.set_margins(left=20, top=20, right=20)
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()
    pdf.add_font("DejaVu", style="", fname=DEJAVU_FONT)
    pdf.add_font("DejaVu", style="B", fname=DEJAVU_BOLD)

    w = pdf.epw  # effective page width po marginesach (~170mm)

    pdf.set_font("DejaVu", style="B", size=13)
    pdf.cell(w, 10, "PROTOKOL ZDAWCZO-ODBIORCZY", align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    pdf.set_font("DejaVu", size=10)

    for raw_line in text.splitlines():
        line = wrap_words(raw_line.strip())
        if not line:
            pdf.ln(3)
            continue
        is_header = line.isupper() or (line.endswith(":") and len(line) < 60)
        if is_header:
            pdf.set_font("DejaVu", style="B", size=10)
        pdf.multi_cell(w, 6, line, new_x="LMARGIN", new_y="NEXT")
        if is_header:
            pdf.set_font("DejaVu", size=10)

    return bytes(pdf.output())


@app.route("/api/export-protocol-pdf", methods=["POST"])
def export_protocol_pdf():
    data = request.get_json()
    if not data or not data.get("protocol_text"):
        return jsonify({"error": "Brak treści protokołu"}), 400
    try:
        pdf_bytes = _build_protocol_pdf(data["protocol_text"])
    except Exception as e:
        return jsonify({"error": f"Błąd generowania PDF: {e}"}), 500

    response = make_response(pdf_bytes)
    response.headers["Content-Type"] = "application/pdf"
    response.headers["Content-Disposition"] = 'attachment; filename="protokol_zdawczo_odbiorczy.pdf"'
    return response


if __name__ == "__main__":
    print("Inicjalizacja bazy RAG...")
    setup_rag()
    print("LeaseGuard uruchomiony na http://localhost:5000")
    app.run(debug=True, host="0.0.0.0", port=5000)
